from pathlib import Path
import ast
import json
import sys

import matplotlib.pyplot as plt
from openpyxl import load_workbook
import yaml as pyyaml

CAMPAIGN_DIR = Path(__file__).resolve().parent
REPO_ROOT = CAMPAIGN_DIR.parents[1]
if str(REPO_ROOT) not in sys.path:
    sys.path.insert(0, str(REPO_ROOT))

from src.design_campaign import generate_design_campaign
from src.design_campaign_executor import WorkflowExecutor
from src.design_campaign_visualization import campaign_parameters_dataframe


def _format_excel_value(value):
    if value is None:
        return ""
    if isinstance(value, (list, tuple, dict)):
        return json.dumps(value)
    return str(value)


def _load_workbook_or_raise(workbook_path, **kwargs):
    try:
        return load_workbook(workbook_path, **kwargs)
    except PermissionError as exc:
        raise PermissionError(
            f"Could not open {workbook_path}. Close the workbook in Excel and rerun "
            "the runtime so it can read and write the design iteration values."
        ) from exc


def _parse_workbook_value(text):
    text = str(text).strip()
    if text == "":
        return None

    try:
        parsed = ast.literal_eval(text)
    except (ValueError, SyntaxError):
        lowered = text.lower()
        if lowered == "true":
            parsed = True
        elif lowered == "false":
            parsed = False
        elif lowered in {"none", "null"}:
            parsed = None
        else:
            parsed = text

    return parsed


def _parameter_type(value):
    if isinstance(value, bool):
        return "bool"
    if isinstance(value, int):
        return "int"
    if isinstance(value, float):
        return "float"
    if isinstance(value, list):
        return "list"
    if isinstance(value, dict):
        return "dict"
    return "string" if value is not None else "float"


def _normalize_header(value):
    return " ".join(str(value or "").strip().lower().split())


def _find_parameter_table_columns(worksheet):
    # Columns A:G are the campaign table. Later columns may contain manual
    # calculations and duplicate headers, so runtime must ignore them.
    campaign_table_last_column = 7
    required_headers = ("component", "parameter", "value", "parameter kind")

    for row in worksheet.iter_rows(max_col=campaign_table_last_column):
        headers = {}
        for cell in row:
            header = _normalize_header(cell.value)
            if header and header not in headers:
                headers[header] = cell.column

        if all(header in headers for header in required_headers):
            return row[0].row, {
                header: headers[header] for header in required_headers
            }

    raise ValueError(
        f"Could not find Component, Parameter, and Value columns in {worksheet.title!r}."
    )


def _get_requirements_worksheet(workbook):
    if "UAV Design Requirements" in workbook.sheetnames:
        return workbook["UAV Design Requirements"]
    return workbook.active


def _load_yaml_data(yaml_path):
    with yaml_path.open("r", encoding="utf-8") as stream:
        return pyyaml.safe_load(stream)


def reset_design_iteration_outputs(workbook_path):
    workbook = _load_workbook_or_raise(workbook_path)
    worksheet = _get_requirements_worksheet(workbook)
    header_row, columns = _find_parameter_table_columns(worksheet)
    cleared = 0

    for row_idx in range(header_row + 1, worksheet.max_row + 1):
        component_id = str(
            worksheet.cell(row_idx, columns["component"]).value or ""
        ).strip()
        parameter_name = str(
            worksheet.cell(row_idx, columns["parameter"]).value or ""
        ).strip()
        parameter_kind = str(
            worksheet.cell(row_idx, columns["parameter kind"]).value or ""
        ).strip().lower()
        if not component_id or not parameter_name or not (
            "comput" in parameter_kind or "optim" in parameter_kind
        ):
            continue

        value_cell = worksheet.cell(row_idx, columns["value"])
        if value_cell.value not in (None, ""):
            value_cell.value = None
            cleared += 1

    workbook.save(workbook_path)
    print(f"Reset {cleared} optimized/computed workbook values in: {workbook_path}")


def _read_workbook_components(workbook_path):
    workbook = _load_workbook_or_raise(workbook_path, data_only=True)
    worksheet = _get_requirements_worksheet(workbook)
    header_row, columns = _find_parameter_table_columns(worksheet)

    components = {}
    for row_idx in range(header_row + 1, worksheet.max_row + 1):
        component_id = str(
            worksheet.cell(row_idx, columns["component"]).value or ""
        ).strip()
        parameter_name = str(
            worksheet.cell(row_idx, columns["parameter"]).value or ""
        ).strip()
        raw_value = worksheet.cell(row_idx, columns["value"]).value
        parameter_kind = str(
            worksheet.cell(row_idx, columns["parameter kind"]).value or ""
        ).strip()
        if not component_id or not parameter_name or not parameter_kind:
            continue

        value = None if raw_value in (None, "") else _parse_workbook_value(raw_value)
        payload = {"type": _parameter_type(value), "kind": parameter_kind}
        if "known" in parameter_kind.lower():
            payload["value"] = value
        components.setdefault(component_id, {})[parameter_name] = payload

    # A sibling '<parameter>_bounds' row supplies bounds for its base parameter,
    # regardless of whether that base is labelled optimization or computed.
    for parameters in components.values():
        for parameter_name, payload in parameters.items():
            bounds_payload = parameters.get(f"{parameter_name}_bounds", {})
            bounds = bounds_payload.get("value")
            if isinstance(bounds, list):
                payload["bounds"] = bounds

    return [
        {"id": component_id, "type": "workbook", "parameters": parameters}
        for component_id, parameters in components.items()
    ]


def load_campaign_from_workbook_and_models(yaml_path, workbook_path):
    workflow_data = _load_yaml_data(yaml_path)
    models = workflow_data.get("models", [])
    if not isinstance(models, list):
        raise ValueError(f"The models section in {yaml_path} must be a list.")

    reset_design_iteration_outputs(workbook_path)
    components = _read_workbook_components(workbook_path)
    parameter_count = sum(len(component["parameters"]) for component in components)
    print(
        f"Created {len(components)} components with {parameter_count} parameters "
        f"from: {workbook_path}"
    )
    return {"components": components, "models": models}


def populate_design_iterations_workbook(workbook_path, campaign):
    values = {}
    for component in campaign.components:
        for parameter in component.optimization_variables + component.computed_variables:
            if parameter.value is None:
                continue
            values[(component.id, parameter.tag)] = parameter.value

    workbook = _load_workbook_or_raise(workbook_path)
    worksheet = _get_requirements_worksheet(workbook)
    header_row, columns = _find_parameter_table_columns(worksheet)

    for row_idx in range(header_row + 1, worksheet.max_row + 1):
        component_id = str(
            worksheet.cell(row_idx, columns["component"]).value or ""
        ).strip()
        parameter_name = str(
            worksheet.cell(row_idx, columns["parameter"]).value or ""
        ).strip()
        key = (component_id, parameter_name)
        if key not in values:
            continue

        worksheet.cell(row_idx, columns["value"]).value = _format_excel_value(values[key])

    workbook.save(workbook_path)

# Select campaign-local design template and workbook.
yaml_path = CAMPAIGN_DIR / "small_battery_UAV_lift_cruise.yaml"
xlsx_template_path = CAMPAIGN_DIR / "UAV Design Iterations.xlsx"

# Generate design campaign object after populating yaml data from workbook values
yaml_data = load_campaign_from_workbook_and_models(yaml_path, xlsx_template_path)
campaign = generate_design_campaign(yaml_data)

executor = WorkflowExecutor(campaign)
results = executor.run()

# If geometry stage returned an airplane object, show it
for r in results:
    obj_key = f"{r.stage_id}__object"
    if obj_key in r.outputs:
        airplane = r.outputs[obj_key]
        airplane.draw_three_view()

# Plot CL vs alpha from vsp_run sweep outputs when available
for r in results:
    sweep_outputs = r.outputs.get(f"{r.stage_id}__sweep_outputs")
    if not isinstance(sweep_outputs, dict):
        continue

    alpha = sweep_outputs.get("Alpha")
    cl = sweep_outputs.get("CL")
    if alpha is None or cl is None or len(alpha) == 0 or len(cl) == 0:
        continue

    n = min(len(alpha), len(cl))
    plt.figure()
    plt.plot(alpha[:n], cl[:n], marker="o")
    plt.xlabel("alpha (deg)")
    plt.ylabel("CL")
    plt.title(f"{r.stage_id}: CL vs alpha")
    plt.grid(True, alpha=0.3)
    plt.tight_layout()
    plt.show()
    break


# Display all the parameters stored in results 
for r in results: 
    print(r.stage_id + " results:/n")
    for key, value in r.outputs.items():
        print(str(key) + ":" + str(value))

parameters_df = campaign_parameters_dataframe(campaign)
print("\nCampaign parameters (grouped by global/component):")
print(parameters_df.to_string(index=False))

populate_design_iterations_workbook(xlsx_template_path, campaign)
print(f"\nUpdated optimized and computed variables in workbook: {xlsx_template_path}")
