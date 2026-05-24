from pathlib import Path
import ast
import json

import matplotlib.pyplot as plt
from openpyxl import load_workbook
import yaml as pyyaml
from examples.small_UAV_lift_cruise.design_campaign import generate_design_campaign
from src.design_campaign_executor import WorkflowExecutor
from src.design_campaign_visualization import campaign_parameters_dataframe


def _format_excel_value(value):
    if value is None:
        return ""
    if isinstance(value, (list, tuple, dict)):
        return json.dumps(value)
    return str(value)


def _load_workbook_or_raise(workbook_path, **kwargs):
    try:
        return load_workbook(workbook_path, **kwargs)
    except PermissionError as exc:
        raise PermissionError(
            f"Could not open {workbook_path}. Close the workbook in Excel and rerun "
            "the runtime so it can read and write the design iteration values."
        ) from exc


def _parse_workbook_value(text, parameter_payload):
    text = str(text).strip()
    if text == "":
        return None

    try:
        parsed = ast.literal_eval(text)
    except (ValueError, SyntaxError):
        lowered = text.lower()
        if lowered == "true":
            parsed = True
        elif lowered == "false":
            parsed = False
        elif lowered in {"none", "null"}:
            parsed = None
        else:
            parsed = text

    parameter_type = str(parameter_payload.get("type", "")).strip().lower()
    if parameter_payload.get("bounds") is not None and isinstance(parsed, list):
        return parsed
    if parameter_type in {"string", "str"} or parsed is None:
        return parsed
    if parameter_type == "bool":
        if isinstance(parsed, str):
            return parsed.strip().lower() in {"1", "true", "yes", "y"}
        return bool(parsed)
    if parameter_type == "int" and not isinstance(parsed, bool):
        return int(float(parsed))
    if parameter_type == "float" and not isinstance(parsed, bool):
        return float(parsed)
    return parsed


def _normalize_header(value):
    return " ".join(str(value or "").strip().lower().split())


def _find_parameter_table_columns(worksheet):
    required_headers = {
        "component": None,
        "parameter": None,
        "value": None,
    }

    for row in worksheet.iter_rows():
        headers = {
            _normalize_header(cell.value): cell.column
            for cell in row
            if str(cell.value or "").strip()
        }
        if not headers:
            continue

        for header in required_headers:
            required_headers[header] = headers.get(header)

        if all(required_headers.values()):
            return row[0].row, required_headers

    raise ValueError(
        f"Could not find Component, Parameter, and Value columns in {worksheet.title!r}."
    )


def _get_requirements_worksheet(workbook):
    if "UAV Design Requirements" in workbook.sheetnames:
        return workbook["UAV Design Requirements"]
    return workbook.active


def _load_yaml_data(yaml_path):
    with yaml_path.open("r", encoding="utf-8") as stream:
        return pyyaml.safe_load(stream)


def _output_parameter_keys(yaml_data):
    output_keys = set()
    for component in yaml_data.get("components", []) or []:
        component_id = str(component.get("id", ""))
        for parameter_name, parameter_payload in (component.get("parameters", {}) or {}).items():
            if not isinstance(parameter_payload, dict):
                continue

            parameter_kind = str(parameter_payload.get("kind", "")).lower()
            if "comput" in parameter_kind or "optim" in parameter_kind:
                output_keys.add((component_id, str(parameter_name)))

    return output_keys


def reset_design_iteration_outputs(workbook_path, yaml_data):
    output_keys = _output_parameter_keys(yaml_data)
    workbook = _load_workbook_or_raise(workbook_path)
    worksheet = _get_requirements_worksheet(workbook)
    header_row, columns = _find_parameter_table_columns(worksheet)
    cleared = 0

    for row_idx in range(header_row + 1, worksheet.max_row + 1):
        component_id = str(
            worksheet.cell(row_idx, columns["component"]).value or ""
        ).strip()
        parameter_name = str(
            worksheet.cell(row_idx, columns["parameter"]).value or ""
        ).strip()
        key = (component_id, parameter_name)
        if key not in output_keys:
            continue

        value_cell = worksheet.cell(row_idx, columns["value"])
        if value_cell.value not in (None, ""):
            value_cell.value = None
            cleared += 1

    workbook.save(workbook_path)
    print(f"Reset {cleared} optimized/computed workbook values in: {workbook_path}")


def _read_workbook_component_values(workbook_path):
    workbook = _load_workbook_or_raise(workbook_path, data_only=True)
    worksheet = _get_requirements_worksheet(workbook)
    header_row, columns = _find_parameter_table_columns(worksheet)

    workbook_values = {}
    for row_idx in range(header_row + 1, worksheet.max_row + 1):
        component_id = str(
            worksheet.cell(row_idx, columns["component"]).value or ""
        ).strip()
        parameter_name = str(
            worksheet.cell(row_idx, columns["parameter"]).value or ""
        ).strip()
        raw_value = worksheet.cell(row_idx, columns["value"]).value
        value_text = str(raw_value or "").strip()
        if component_id and parameter_name and value_text != "":
            workbook_values[(component_id, parameter_name)] = value_text

    return workbook_values


def load_yaml_with_workbook_parameters(yaml_path, workbook_path):
    data = _load_yaml_data(yaml_path)
    reset_design_iteration_outputs(workbook_path, data)

    workbook_values = _read_workbook_component_values(workbook_path)
    updates = 0

    for component in data.get("components", []) or []:
        component_id = str(component.get("id", ""))
        parameters = component.get("parameters", {}) or {}
        for parameter_name, parameter_payload in parameters.items():
            key = (component_id, str(parameter_name))
            if key not in workbook_values or not isinstance(parameter_payload, dict):
                continue

            parameter_kind = str(parameter_payload.get("kind", "")).lower()
            if "comput" in parameter_kind or "optim" in parameter_kind:
                continue

            parsed_value = _parse_workbook_value(workbook_values[key], parameter_payload)
            if parsed_value is None:
                continue

            if parameter_payload.get("bounds") is not None and isinstance(parsed_value, list):
                parameter_payload["bounds"] = parsed_value
            else:
                parameter_payload["value"] = parsed_value
            updates += 1

    print(f"Loaded {updates} component parameter values from: {workbook_path}")
    return data


def populate_design_iterations_workbook(workbook_path, campaign):
    values = {}
    for component in campaign.components:
        for parameter in component.optimization_variables + component.computed_variables:
            if parameter.value is None:
                continue
            values[(component.id, parameter.tag)] = parameter.value

    workbook = _load_workbook_or_raise(workbook_path)
    worksheet = _get_requirements_worksheet(workbook)
    header_row, columns = _find_parameter_table_columns(worksheet)

    for row_idx in range(header_row + 1, worksheet.max_row + 1):
        component_id = str(
            worksheet.cell(row_idx, columns["component"]).value or ""
        ).strip()
        parameter_name = str(
            worksheet.cell(row_idx, columns["parameter"]).value or ""
        ).strip()
        key = (component_id, parameter_name)
        if key not in values:
            continue

        worksheet.cell(row_idx, columns["value"]).value = _format_excel_value(values[key])

    workbook.save(workbook_path)

# Select design template
repo_root = Path(__file__).resolve().parents[2]
yaml_path = repo_root / "templates" / "aerospace" / "small_H2_UAV_lift_cruise.yaml"
xlsx_template_path = Path(__file__).with_name("UAV Design Iterations.xlsx")

# Generate design campaign object after populating yaml data from workbook values
yaml_data = load_yaml_with_workbook_parameters(yaml_path, xlsx_template_path)
campaign = generate_design_campaign(yaml_data)

executor = WorkflowExecutor(campaign)
results = executor.run()

# If geometry stage returned an airplane object, show it
for r in results:
    obj_key = f"{r.stage_id}__object"
    if obj_key in r.outputs:
        airplane = r.outputs[obj_key]
        airplane.draw_three_view()

# Plot CL vs alpha from vsp_run sweep outputs when available
for r in results:
    sweep_outputs = r.outputs.get(f"{r.stage_id}__sweep_outputs")
    if not isinstance(sweep_outputs, dict):
        continue

    alpha = sweep_outputs.get("Alpha")
    cl = sweep_outputs.get("CL")
    if alpha is None or cl is None or len(alpha) == 0 or len(cl) == 0:
        continue

    n = min(len(alpha), len(cl))
    plt.figure()
    plt.plot(alpha[:n], cl[:n], marker="o")
    plt.xlabel("alpha (deg)")
    plt.ylabel("CL")
    plt.title(f"{r.stage_id}: CL vs alpha")
    plt.grid(True, alpha=0.3)
    plt.tight_layout()
    plt.show()
    break


# Display all the parameters stored in results 
for r in results: 
    print(r.stage_id + " results:/n")
    for key, value in r.outputs.items():
        print(str(key) + ":" + str(value))

parameters_df = campaign_parameters_dataframe(campaign)
print("\nCampaign parameters (grouped by global/component):")
print(parameters_df.to_string(index=False))

csv_path = Path(__file__).with_name("campaign_iterations.csv")
parameters_df.to_csv(csv_path, index=False)
print(f"\nSaved parameters dataframe to: {csv_path}")

populate_design_iterations_workbook(xlsx_template_path, campaign)
print(f"\nUpdated optimized and computed variables in workbook: {xlsx_template_path}")
